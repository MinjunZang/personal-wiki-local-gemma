"""Read original files from vault/raw/ into located sections.

This module only reads. Each section keeps a human-checkable location:
  .docx -> heading + paragraph numbers (¶), counted like python-docx / Word paragraphs
  .pptx -> slide number and slide title
  .md / .txt / Word-HTML .doc -> Markdown heading path
"""
import hashlib
import html
import re
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pptx

SUPPORTED = {".docx", ".pptx", ".doc", ".md", ".txt", ".html", ".htm"}


@dataclass
class Section:
    label: str                      # heading, slide title, or heading path
    unit: str | None                # "¶", "slide", or None (no numeric position)
    lines: list[tuple[int, str]] = field(default_factory=list)  # (position, text)

    @property
    def text(self) -> str:
        return "\n".join(text for _, text in self.lines)

    @property
    def location(self) -> str:
        return format_location(self.label, self.unit, self.lines)


def format_location(label: str, unit: str | None, lines: list[tuple[int, str]]) -> str:
    if unit is None or not lines:
        return label
    first, last = lines[0][0], lines[-1][0]
    if unit == "slide":
        return label if first == last else f"{label} (slides {first}–{last})"
    span = f"{unit}{first}" if first == last else f"{unit}{first}–{last}"
    return f"{label} ({span})"


def is_source_file(path: Path) -> bool:
    return (path.is_file() and path.suffix.lower() in SUPPORTED
            and not path.name.startswith(("~$", ".")))


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_sections(path: Path) -> list[Section]:
    ext = path.suffix.lower()
    if ext == ".docx":
        return _docx_sections(path)
    if ext == ".pptx":
        return _pptx_sections(path)
    raw = path.read_bytes()
    if ext == ".doc" and raw.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValueError(f"{path.name} is a binary Word 97 file; save it as .docx first.")
    text = raw.decode("utf-8", errors="replace")
    if ext in {".doc", ".html", ".htm"}:
        text = _html_to_text(text)
    return markdown_sections(text, fallback_title=path.stem)


def _html_to_text(markup: str) -> str:
    # Word/LectMate exports wrap Markdown in HTML and use <br /> for newlines.
    markup = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>", "\n", markup)
    markup = re.sub(r"<[^>]+>", "", markup)
    return html.unescape(markup)


def _docx_heading(paragraph, text: str) -> bool:
    """The sources use the 'Normal' style everywhere, so headings are detected by shape:
    a short line that is not a list item, not a question, and not a sentence."""
    is_list_item = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
    return (not is_list_item and 3 <= len(text) <= 60
            and not text.endswith(("?", "？", ".", "。", ":", "："))
            and not re.fullmatch(r"\d+\.", text))


def _docx_sections(path: Path) -> list[Section]:
    document = docx.Document(str(path))
    sections = [Section(label="Start", unit="¶")]
    pending_number = ""  # e.g. a lone "3." paragraph followed by its question
    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        if re.fullmatch(r"\d+\.", text):
            pending_number = text + " "
            continue
        if _docx_heading(paragraph, text) and not pending_number:
            sections.append(Section(label=text, unit="¶", lines=[(number, f"## {text}")]))
            continue
        pPr = paragraph._p.pPr
        if pPr is not None and pPr.numPr is not None:
            level = pPr.numPr.ilvl.val if pPr.numPr.ilvl is not None else 0
            text = "  " * level + "- " + text
        sections[-1].lines.append((number, pending_number + text))
        pending_number = ""
    return [s for s in sections if s.lines]


def _pptx_sections(path: Path) -> list[Section]:
    deck = pptx.Presentation(str(path))
    sections = []
    for number, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    texts.append(" | ".join(c for c in cells if c))
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        title = title or (texts[0] if texts else "Untitled")
        sections.append(Section(
            label=f"Slide {number}: {title}", unit="slide",
            lines=[(number, t) for t in texts],
        ))
    return sections


def markdown_sections(text: str, fallback_title: str) -> list[Section]:
    sections = [Section(label=fallback_title, unit=None)]
    path: list[str] = []  # heading stack
    for line in text.splitlines():
        line = line.rstrip()
        match = re.match(r"^(#{1,6})\s+(.*)", line)
        if match:
            level, heading = len(match.group(1)), match.group(2).strip()
            path = path[:level - 1] + [heading]
            # Keep the document title out of the location path.
            label = " > ".join(path[1:]) if len(path) > 1 else path[0]
            sections.append(Section(label=label, unit=None, lines=[(0, line)]))
        elif line.strip():
            sections[-1].lines.append((0, line))
    return [s for s in sections if s.lines]
